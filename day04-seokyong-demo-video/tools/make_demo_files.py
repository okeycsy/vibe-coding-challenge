#!/usr/bin/env python3
"""시연용 실적 파일 4종을 만든다 (모두 가상 데이터, 값은 고정 시드라 매번 같다).

사용: python3 make_demo_files.py [출력 폴더=../demo-files]

기관명·기관코드는 앱의 mockData.js(INSTITUTIONS)에서 읽으므로 앱에 등록된 기관만 쓴다.
출력 폴더는 저장소에 넣지 않는다(.gitignore). 실제 테스트 실적 원본이 있으면 같은 이름으로 이 폴더에 두면 된다.

  오픈뱅킹 테스트 실적.xlsx   기준월 | 기관명 | API 종류별 건수 5열 (50개 기관 × 3개월 = 150행, 날짜 열)
  CD공동망 테스트 실적.xlsx   제목 행 "2026년 3분기 …" + 기관코드 | 금융기관명 | 이용실적 (50곳 + 합계, 표 제목에서 분기를 읽음)
  타행환 테스트 실적.xlsx     같은 구조 (40곳 + 합계)
  기업회계 테스트 실적.xlsx   기관코드 | 기관명 | 2026년 4월 … 2026년 9월 (35곳, 날짜 머리글 → 3분기 밖 4~6월은 제외됨)
"""
import os
import random
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
MOCK = os.path.join(HERE, '..', '..', 'day01-gyeongsik-fee-calculator', 'src', 'data', 'mockData.js')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'demo-files')


def institutions():
    src = open(MOCK, encoding='utf-8').read()
    block = re.search(r'export const INSTITUTIONS = \[(.*?)\n\];', src, re.S).group(1)
    items = re.findall(r"\{ id: '([^']+)', name: '([^']+)', code: '([^']+)'", block)
    return [(code, name) for _id, name, code in items]


def header_style(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E8EEF9')
        cell.alignment = Alignment(horizontal='center')


def widths(ws, cols):
    for i, w in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def write_openbanking(insts, rng):
    t = [i for i in insts if i[0] in ('T001', 'T002', 'T004', 'T008')]  # 한빛·미래온·다온·가온
    o = [i for i in insts if i[0].startswith('O')]
    rows = t + o  # 50곳
    assert len(rows) == 50, len(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '월별_API_원본'
    ws.append(['기준월', '기관명', '거래내역조회', '계좌실명조회', '잔액조회', '출금이체', '입금이체'])
    header_style(ws, 1, 7)
    scale = {code: rng.uniform(0.25, 1.0) ** 2 for code, _ in rows}
    scale['T001'] = 1.0  # 한빛은행이 가장 크게
    scale['O001'] = 1.6  # 우체국
    for month in ('2026-07', '2026-08', '2026-09'):
        for code, name in rows:
            s = scale[code]
            ws.append([month, name,
                       int(s * rng.uniform(9e6, 14e6)), int(s * rng.uniform(2.5e6, 4.5e6)), int(s * rng.uniform(6e6, 9e6)),
                       int(s * rng.uniform(3.5e6, 5.5e6)), int(s * rng.uniform(2.5e6, 4.5e6))])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=7):
        for cell in row:
            cell.number_format = '#,##0'
    widths(ws, [11, 20, 14, 14, 12, 12, 12])
    wb.save(os.path.join(OUT, '오픈뱅킹 테스트 실적.xlsx'))


def write_quarter_table(insts, rng, name, title, count, lo, hi, override=None):
    rows = [i for i in insts if i[0].startswith('T')][:count]
    override = override or {}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '전분기 이용실적'
    ws.append([title])
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    ws.append(['기관코드', '금융기관명', '이용실적'])
    header_style(ws, 2, 3)
    total = 0
    for code, inst in rows:
        v = int(rng.uniform(lo, hi))
        v = override.get(code, v)
        total += v
        ws.append([code, inst, v])
    ws.append(['', '합계', total])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    for row in ws.iter_rows(min_row=3, min_col=3, max_col=3):
        row[0].number_format = '#,##0'
    widths(ws, [11, 20, 14])
    wb.save(os.path.join(OUT, name))


def write_fininfo(insts, rng):
    rows = [i for i in insts if i[0].startswith('T')][:35]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '월별 전문 건수'
    months = [f'2026년 {m}월' for m in range(4, 10)]
    ws.append(['기관코드', '기관명', *months])
    header_style(ws, 1, 2 + len(months))
    for code, inst in rows:
        base = rng.uniform(5e3, 4e4)
        ws.append([code, inst, *[int(base * rng.uniform(0.85, 1.15)) for _ in months]])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=2 + len(months)):
        for cell in row:
            cell.number_format = '#,##0'
    widths(ws, [11, 20] + [13] * len(months))
    wb.save(os.path.join(OUT, '기업회계 테스트 실적.xlsx'))


def main():
    os.makedirs(OUT, exist_ok=True)
    insts = institutions()
    rng = random.Random(20261027)
    write_openbanking(insts, rng)
    # 한빛은행(첫 기관) 회비의 1원 단위가 5 근처가 되는 값: 4번 특이사항(10원 절상·차이분 절삭) 보정 배지가 화면에 보이도록
    write_quarter_table(insts, rng, 'CD공동망 테스트 실적.xlsx', '2026년 3분기 CD공동망 이용실적', 50, 3.2e5, 9.5e5, override={'T001': 589328})
    write_quarter_table(insts, rng, '타행환 테스트 실적.xlsx', '2026년 3분기 타행환 이용실적', 40, 2.0e5, 7.5e5)
    write_fininfo(insts, rng)
    for f in sorted(os.listdir(OUT)):
        print('wrote', os.path.join(OUT, f))


if __name__ == '__main__':
    main()
